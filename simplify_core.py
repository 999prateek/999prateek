import os
import re
import sqlite3
import hashlib
from pathlib import Path
from datetime import datetime

def db_connect(db_path):
    Path(db_path).parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(db_path, check_same_thread=False)
    con.row_factory = sqlite3.Row
    return con

def init_db(db_path):
    with db_connect(db_path) as con:
        con.executescript("""
        CREATE TABLE IF NOT EXISTS sessions (
          session_id TEXT PRIMARY KEY, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS messages (
          id INTEGER PRIMARY KEY AUTOINCREMENT, session_id TEXT NOT NULL,
          role TEXT NOT NULL, content TEXT NOT NULL, language TEXT,
          sources TEXT, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS documents (
          id INTEGER PRIMARY KEY AUTOINCREMENT, filename TEXT NOT NULL,
          stored_path TEXT NOT NULL, sha256 TEXT NOT NULL UNIQUE,
          chunks INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS chunks (
          id INTEGER PRIMARY KEY AUTOINCREMENT, document_id INTEGER NOT NULL,
          chunk_index INTEGER NOT NULL, content TEXT NOT NULL,
          FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
        );
        """)

def detect_language(text):
    if re.search(r'[\u0900-\u097F]', text): return "Hindi / Devanagari"
    if re.search(r'[\u0A80-\u0AFF]', text): return "Gujarati"
    if re.search(r'[\u0980-\u09FF]', text): return "Bengali"
    if re.search(r'[\u0B80-\u0BFF]', text): return "Tamil"
    if re.search(r'[\u0C00-\u0C7F]', text): return "Telugu"
    if re.search(r'[\u0C80-\u0CFF]', text): return "Kannada"
    if re.search(r'[\u0D00-\u0D7F]', text): return "Malayalam"
    words=set(re.findall(r"[a-z]+",text.lower()))
    if words & {"hai","hain","kya","kaise","mujhe","chahiye","batao","nahi","mera","mein","ka","ki","ke"}:
        return "Hinglish"
    return "English / Latin-script"

def save_message(db_path, session_id, role, content, language, sources):
    import json
    now=datetime.utcnow().isoformat(timespec="seconds")+"Z"
    with db_connect(db_path) as con:
        con.execute("INSERT OR IGNORE INTO sessions(session_id,created_at) VALUES (?,?)",(session_id,now))
        con.execute("INSERT INTO messages(session_id,role,content,language,sources,created_at) VALUES (?,?,?,?,?,?)",
                    (session_id,role,content,language,json.dumps(sources or [],ensure_ascii=False),now))

def get_history(db_path, session_id):
    import json
    with db_connect(db_path) as con:
        rows=con.execute("SELECT role,content,sources FROM messages WHERE session_id=? ORDER BY id",(session_id,)).fetchall()
    result=[]
    for r in rows:
        try: sources=json.loads(r["sources"] or "[]")
        except Exception: sources=[]
        result.append({"role":r["role"],"content":r["content"],"sources":sources})
    return result

def split_chunks(text, size=1000, overlap=150):
    text=re.sub(r'\s+',' ',text).strip()
    result=[]
    start=0
    while start < len(text):
        end=min(len(text),start+size)
        result.append(text[start:end])
        if end == len(text): break
        start=max(start+1,end-overlap)
    return result

def extract_file(path):
    suffix=Path(path).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader=PdfReader(str(path))
        pages=[]
        for i,p in enumerate(reader.pages):
            text=p.extract_text() or ""
            if text.strip(): pages.append(f"[Page {i+1}] {text}")
        return "\n".join(pages)
    if suffix == ".docx":
        from docx import Document
        doc=Document(str(path))
        parts=[p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if suffix in (".txt",".md",".csv"):
        return Path(path).read_text(encoding="utf-8",errors="replace")
    raise ValueError("Unsupported file format")

def ingest_file(db_path, path, original_name):
    content=extract_file(path)
    chunks=split_chunks(content)
    digest=hashlib.sha256(Path(path).read_bytes()).hexdigest()
    with db_connect(db_path) as con:
        existing=con.execute("SELECT id FROM documents WHERE sha256=?",(digest,)).fetchone()
        if existing: return 0
        cur=con.execute("INSERT INTO documents(filename,stored_path,sha256,chunks,created_at) VALUES (?,?,?,?,?)",
                        (original_name,str(path),digest,len(chunks),datetime.utcnow().isoformat(timespec="seconds")+"Z"))
        doc_id=cur.lastrowid
        con.executemany("INSERT INTO chunks(document_id,chunk_index,content) VALUES (?,?,?)",
                        [(doc_id,i,c) for i,c in enumerate(chunks)])
    return len(chunks)

def retrieve_docs(db_path, question, limit=5):
    stop=set("the a an is are was were to of and or in on for with how what when where can i you me my please tell about this that from explain summarize document file".split())
    query=set(re.findall(r'[\w\u0900-\u0D7F]+',question.lower()))-stop
    if not query: return []
    with db_connect(db_path) as con:
        rows=con.execute("SELECT c.content,d.filename,c.chunk_index FROM chunks c JOIN documents d ON d.id=c.document_id").fetchall()
    scored=[]
    for r in rows:
        words=set(re.findall(r'[\w\u0900-\u0D7F]+',r["content"].lower()))
        score=len(query & words)/max(1,len(query))
        if score > 0:
            scored.append({"score":score,"filename":r["filename"],"chunk_index":r["chunk_index"],"content":r["content"]})
    scored.sort(key=lambda x:x["score"],reverse=True)
    return scored[:limit]

def web_search(query, limit=5):
    try:
        from ddgs import DDGS
        results=[]
        with DDGS() as ddgs:
            for r in ddgs.text(query, max_results=limit):
                results.append({"title":r.get("title","Web result"),"url":r.get("href") or r.get("url",""),"snippet":r.get("body","")})
        return results
    except Exception:
        return []

def call_llm(question, context, web_results, target_language, history):
    endpoint=os.getenv("OPENAI_BASE_URL","").strip()
    api_key=os.getenv("OPENAI_API_KEY","").strip()
    model=os.getenv("LLM_MODEL","")
    if not endpoint or not api_key or not model:
        return None
    try:
        import requests
        web_context="\n\n".join([f"Title: {x['title']}\nURL: {x['url']}\nSnippet: {x['snippet']}" for x in web_results])
        messages=[
            {"role":"system","content":(
                "You are Simplify AI, a helpful general-purpose assistant. Answer questions across domains. "
                "Use the user's uploaded document excerpts when relevant, and use web search snippets for current facts. "
                "Do not pretend you accessed a source you did not access. If evidence conflicts, explain the conflict. "
                "For high-stakes medical, legal, financial, or safety questions, give careful general information and recommend appropriate expert/official verification. "
                "If context is insufficient, say what is unknown instead of inventing details. "
                f"Respond in {target_language}; if Auto-detect, match the user's language. "
                "Keep answers clear, detailed enough to be useful, and cite sources by title/URL when available."
            )},
        ]
        for item in history[-8:]:
            if item["role"] in ("user","assistant"):
                messages.append({"role":item["role"],"content":item["content"][:3000]})
        user_prompt=f"""QUESTION:
{question}

UPLOADED DOCUMENT EXCERPTS:
{context if context else '[No relevant uploaded-document excerpts retrieved]'}

LIVE WEB SEARCH RESULTS:
{web_context if web_context else '[No web search results available]'}

Instructions:
- Answer the question directly.
- Prefer uploaded-document evidence when the user asks about their file, but clearly distinguish it from general knowledge.
- If a document does not contain the answer, say so; do not invent a document-based answer.
- For time-sensitive claims, use web snippets cautiously and include source URLs where relevant.
"""
        messages.append({"role":"user","content":user_prompt})
        response=requests.post(endpoint.rstrip("/")+"/chat/completions",
            headers={"Authorization":f"Bearer {api_key}","Content-Type":"application/json"},
            json={"model":model,"messages":messages,"temperature":0.25},
            timeout=60)
        response.raise_for_status()
        return response.json()["choices"][0]["message"]["content"].strip()
    except Exception:
        return None

def answer_question(db_path, question, response_language, live_search, use_docs, session_id):
    language=response_language
    if language == "Auto-detect":
        language=detect_language(question)
    doc_hits=retrieve_docs(db_path,question) if use_docs else []
    web_results=web_search(question) if live_search else []
    context="\n\n".join([f"Document: {h['filename']} (chunk {h['chunk_index']+1})\n{h['content']}" for h in doc_hits])
    with db_connect(db_path) as con:
        rows=con.execute("SELECT role,content FROM messages WHERE session_id=? ORDER BY id DESC LIMIT 10",(session_id,)).fetchall()
    history=[{"role":r["role"],"content":r["content"]} for r in reversed(rows)]
    answer=call_llm(question,context,web_results,language,history)
    sources=[]
    for h in doc_hits:
        label=f"Uploaded document: {h['filename']} (chunk {h['chunk_index']+1})"
        if label not in sources: sources.append(label)
    for w in web_results:
        if w.get("url"): sources.append({"title":w.get("title","Web result"),"url":w["url"]})
    if answer is None:
        if doc_hits:
            excerpts="\n\n".join([f"**{h['filename']}**\n> {h['content'][:650]}{'…' if len(h['content'])>650 else ''}" for h in doc_hits[:3]])
            answer=("I found relevant excerpts in your uploaded files, but the AI model is not configured or could not be reached. "
                    "Here are the matching passages; configure an LLM endpoint to get a synthesized answer.\n\n"+excerpts)
        elif web_results:
            answer=("I found these public web results, but the AI model is not configured or could not be reached. "
                    "Open the links to review the information:\n\n"+
                    "\n\n".join([f"- **[{w['title']}]({w['url']})** — {w['snippet']}" for w in web_results[:5]]))
        else:
            answer=("Simplify's AI model is not configured yet, so I can't generate a full answer. "
                    "Set `OPENAI_BASE_URL`, `OPENAI_API_KEY`, and `LLM_MODEL` in the hosting environment. "
                    "You can still upload files; relevant excerpts will be shown when retrieval finds a match.")
    return answer, sources, {"language":language}

def get_documents(db_path):
    with db_connect(db_path) as con:
        rows=con.execute("SELECT id,filename,stored_path,chunks,created_at FROM documents ORDER BY id DESC").fetchall()
    return [dict(r) for r in rows]

def delete_document(db_path, doc_id):
    with db_connect(db_path) as con:
        con.execute("DELETE FROM chunks WHERE document_id=?",(doc_id,))
        con.execute("DELETE FROM documents WHERE id=?",(doc_id,))

def get_analytics(db_path):
    import pandas as pd
    with db_connect(db_path) as con:
        messages=con.execute("SELECT COUNT(*) FROM messages").fetchone()[0]
        sessions=con.execute("SELECT COUNT(DISTINCT session_id) FROM messages").fetchone()[0]
        documents=con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        lang=con.execute("SELECT COALESCE(language,'Unknown') language,COUNT(*) count FROM messages WHERE role='user' GROUP BY language ORDER BY count DESC").fetchall()
        recent=con.execute("SELECT role,substr(content,1,130) preview,language,created_at FROM messages ORDER BY id DESC LIMIT 20").fetchall()
    return {"messages":messages,"sessions":sessions,"documents":documents,
            "languages":pd.DataFrame([dict(r) for r in lang]),"recent":pd.DataFrame([dict(r) for r in recent])}

def transcribe_audio(audio_bytes):
    try:
        from faster_whisper import WhisperModel
        import tempfile
        model=WhisperModel(os.getenv("WHISPER_MODEL","small"),device="cpu",compute_type="int8")
        with tempfile.NamedTemporaryFile(suffix=".wav") as f:
            f.write(audio_bytes); f.flush()
            segments,info=model.transcribe(f.name,beam_size=3,vad_filter=True)
            return " ".join(s.text.strip() for s in segments).strip(),f"Detected {info.language}."
    except Exception:
        return "", "Voice transcription is optional. Install faster-whisper and ffmpeg or connect an ASR provider."

def synthesize_audio(text, language):
    try:
        import asyncio, edge_tts, tempfile
        voices={"Hindi":"hi-IN-SwaraNeural","Hinglish":"hi-IN-SwaraNeural","English":"en-IN-NeerjaNeural","Auto-detect":"en-IN-NeerjaNeural"}
        voice=voices.get(language,"hi-IN-SwaraNeural")
        async def run():
            with tempfile.NamedTemporaryFile(suffix=".mp3",delete=False) as f: path=f.name
            await edge_tts.Communicate(text[:3500],voice).save(path)
            data=Path(path).read_bytes()
            Path(path).unlink(missing_ok=True)
            return data
        return asyncio.run(run()),"Audio generated."
    except Exception:
        return None,"Text-to-speech is optional. Install edge-tts or configure an Indic TTS provider."
